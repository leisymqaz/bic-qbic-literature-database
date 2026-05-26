#!/usr/bin/env python
"""Build a Chinese report on common BIC data types and studied structures."""

from __future__ import annotations

import argparse
import sqlite3
from pathlib import Path

from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH


PROJECT_DIR = Path(__file__).resolve().parents[1]
DEFAULT_DB = PROJECT_DIR / "data" / "bic_literature_verified.sqlite"
DEFAULT_DOCX = PROJECT_DIR / "outputs" / "BIC_common_data_and_structures_report.docx"
DEFAULT_MD = PROJECT_DIR / "reports" / "BIC_common_data_and_structures_report.md"


COMMON_DATA = [
    ["结构与材料", "周期、半径、孔径、厚度、高度、gap、非对称参数 alpha、阵列尺寸 N、基底/层叠材料、折射率 n/k", "结构示意图、SEM/显微图、unit cell 参数表、材料层叠图", "建立 device/material/geometry_params 表；后续生成 eps2d 和 FFT 通道"],
    ["光谱响应", "T/R/A 光谱、PL/散射谱、反射谱、角分辨谱、线宽/FWHM、峰/谷位置", "transmission/reflection spectrum、angle-resolved spectrum、Fano line-shape plot", "提取 lambda/frequency、linewidth、Q、Fano q；图谱数字化入 HDF5"],
    ["Q 与 qBIC 标量标签", "Q factor、logQ、radiative/nonradiative Q、Q~alpha^-2、Q 饱和、fit error", "Q-alpha 曲线、Q vs array size、Q vs perturbation、Fano fitting 图", "后续做 scaling audit 和 loss decomposition"],
    ["Fano/TCMT 参数", "Fano q、gamma、omega0/lambda0、background、coupling coefficient、direct/resonant channel", "Fano fit overlay、TCMT schematic、pole-zero/Fano curve", "用于 Fourier harmonics -> radiative coupling 因果研究"],
    ["k-space / 能带", "band structure、dispersion、light cone、Gamma 点、off-Gamma BIC、BZ folding、topological charge", "band diagram、k-space map、angle-resolved spectra、momentum-space plot", "连接 BIC 机制、拓扑 charge 与辐射通道"],
    ["偏振与拓扑", "polarization vortex、far-field polarization、Stokes parameters、circular dichroism、chirality、topological charge", "polarization map、vortex field、CD spectrum、Poincare index 图", "判断 topological BIC、chiral BIC、Janus/channel-selective BIC"],
    ["近场/模式分布", "E/H field、multipole decomposition、LDOS/PLDOS、mode symmetry、field enhancement", "near-field map、multipole bar chart、LDOS map、mode profile", "解释 high-Q 来源、非线性/传感增强、辐射泄漏"],
    ["应用性能", "sensing FOM、lasing threshold、nonreciprocal contrast、nonlinear conversion、wavefront phase、holography bandwidth", "传感曲线、laser spectrum、phase map、hologram result、nonlinear response plot", "作为二级标签，不应替代基础物理标签"],
]


STRUCTURES = [
    ["Photonic crystal slab / PhC slab", "孔阵列、周期 slab、few-unit-cell slab、twisted/moire PhC", "symmetry-protected/accidental/topological BIC；band structure；light cone；finite-size quasi-BIC；polarization vortex", "能带、角分辨谱、Q、lambda、topological charge、模式场、有限尺寸 Q"],
    ["Diffraction/subwavelength grating", "1D/2D grating、dielectric grating、guided-mode resonance grating", "leaky mode、guided-mode resonance、symmetry-protected mode、normal-incidence filter", "透射/反射谱、Q、Fano 线型、period、thickness、band/leaky-mode 数据"],
    ["All-dielectric metasurface", "Si/TiO2/GaAs/SiN nanodisk、nanobar、nanohole、split/slot unit cell、asymmetric hole disk", "qBIC/Fano、非对称参数调谐、高 Q、传感、SERS/非线性、field enhancement", "结构参数、材料、T/R/A、Q、lambda、Fano q、near-field、multipole、alpha scaling"],
    ["Bilayer / moire / superlattice metasurface", "双层互补 metasurface、moire/twisted PhC、supercell/BZ folding", "BZ folding BIC、twist-induced vortex、mode hybridization、degenerate BIC", "twist angle、superperiod、band/k-space、topological charge、OAM/vortex、角分辨谱"],
    ["Dielectric nanoparticle / nanosphere / rod arrays", "sphere chain、rod circular array、nanoparticle arrays", "light guiding above light line、finite-array qBIC、OAM/near-BIC、collective resonances", "阵列尺寸、粒子半径/间距、Q、guided resonance、field enhancement、finite-size scaling"],
    ["Waveguide arrays / Lieb lattice", "optical waveguide array、Lieb photonic lattice、compact localized state", "早期实验 BIC、Fano compact state、flat-band/localized state", "模式分布、传播图、谱/耦合参数、continuum 中 localized state 证据"],
    ["Active/lasing BIC arrays", "cylindrical nanoresonator array、GaAs nanoantenna array、perovskite/vortex microlaser", "BIC laser、directional lasing、threshold、far-field/vortex emission", "lasing wavelength、threshold、Q、far-field、polarization/OAM、pump-dependent spectra"],
    ["Chiral / Janus / polarization BIC metasurface", "chiral silicon metasurface、spin-orbit locking、Janus upward/downward channels", "chiral BIC、circular dichroism、spin/orbit locking、channel-dependent topological charge", "CD spectra、Stokes/polarization map、topological charge、up/down radiation, Fano response"],
    ["High-Q guided-mode / membrane / multilayer resonators", "photoresist perturbation on SOI、PMMA/SiN/SiO2 membrane、air-mode metasurface", "ultrahigh-Q GMR、million-Q resonator、air-confined qBIC、vibrational strong coupling", "layer stack、period/hole size、Q up to 1e6 scale、momentum-space spectroscopy、FWHM、Q-alpha scaling"],
]


def add_table(doc: Document, headers, rows):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    hdr = table.rows[0].cells
    for i, h in enumerate(headers):
        hdr[i].text = h
        for r in hdr[i].paragraphs[0].runs:
            r.bold = True
            r.font.color.rgb = RGBColor(255, 255, 255)
        shade_cell(hdr[i], "1F4E78")
    for row in rows:
        cells = table.add_row().cells
        for i, val in enumerate(row):
            cells[i].text = str(val)
            cells[i].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP
    doc.add_paragraph()


def shade_cell(cell, fill):
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    cell._tc.get_or_add_tcPr().append(shd)


def build_text(conn: sqlite3.Connection) -> str:
    obs_counts = conn.execute("SELECT observation_type, COUNT(*) FROM physical_observations GROUP BY observation_type ORDER BY COUNT(*) DESC").fetchall()
    lines = [
        "# BIC/qBIC 文献常见数据类型与结构研究对象报告",
        "",
        "## 1. 结论概览",
        "",
        "BIC/qBIC 文献中最常见的数据不是单一的 Q 值，而是一组互相关联的物理证据：结构参数、材料、光谱、Q/linewidth/Fano 拟合、k-space/能带、偏振拓扑、近场模式和应用性能。后续数据库应以 `物理观测-证据链` 为核心，而不是只按论文标题或 BIC 类型分类。",
        "",
        "当前数据库已从本地 PDF 和部分联网网页中抽取 source-quoted 物理观测候选，按类型统计如下：",
        "",
        "| 物理观测类型 | 当前候选条数 |",
        "|---|---:|",
    ]
    for t, n in obs_counts:
        lines.append(f"| {t} | {n} |")
    lines.extend(["", "## 2. BIC 文献常见物理量和图", "", "| 数据类别 | 常见物理量 | 常见图/表 | 数据库用途 |", "|---|---|---|---|"])
    for row in COMMON_DATA:
        lines.append("| " + " | ".join(row) + " |")
    lines.extend(["", "## 3. 常见结构与研究内容矩阵", "", "| 结构类型 | 典型结构 | 主要研究方面 | 常见可入库数据 |", "|---|---|---|---|"])
    for row in STRUCTURES:
        lines.append("| " + " | ".join(row) + " |")
    lines.extend(
        [
            "",
            "## 4. 入库优先级建议",
            "",
            "1. 第一优先级：结构参数、材料、光谱坐标、Q、lambda/frequency、linewidth/Fano q，因为它们可直接支持横向比较和后续建模。",
            "2. 第二优先级：k-space、band structure、polarization vortex/topological charge，因为它们决定 BIC 机制和新 idea 的物理深度。",
            "3. 第三优先级：near-field、multipole、LDOS/PLDOS、field enhancement，用于解释应用性能和辐射通道。",
            "4. 应用指标如 sensing FOM、lasing threshold、nonlinear conversion 应作为派生/应用层标签，不应替代基础物理标签。",
            "",
            "## 5. 审阅注意",
            "",
            "当前 `physical_observations` 多数仍是候选行。它们来源于 PDF 文本或网页 quote，适合人工审阅，但只有经过 page/figure/table/source 核验后才能作为最终事实。图谱中的 Q、Fano q、linewidth 应优先寻找原始数据或补充材料；如果只能从图中数字化，需要记录坐标轴校准、图号、digitization uncertainty 和拟合公式。",
        ]
    )
    return "\n".join(lines)


def build_docx(md_text: str, out: Path) -> None:
    doc = Document()
    sec = doc.sections[0]
    sec.top_margin = Inches(0.8)
    sec.bottom_margin = Inches(0.8)
    sec.left_margin = Inches(0.75)
    sec.right_margin = Inches(0.75)
    doc.styles["Normal"].font.name = "Microsoft YaHei"
    doc.styles["Normal"].font.size = Pt(10.5)
    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run("BIC/qBIC 文献常见数据类型与结构研究对象报告")
    run.bold = True
    run.font.size = Pt(17)
    run.font.color.rgb = RGBColor(31, 78, 121)
    doc.add_paragraph("本报告用于回答：BIC 文献中常见的数据有哪些物理量或图？这些文献研究了哪些结构，以及围绕这些结构记录了哪些数据。")
    doc.add_heading("1. 结论概览", level=1)
    doc.add_paragraph("BIC/qBIC 文献的核心数据通常由结构参数、材料、光谱、Q/Fano/linewidth、k-space/能带、偏振拓扑和近场模式共同组成。数据库应以物理观测和证据链为中心。")
    doc.add_heading("2. 常见物理量和图", level=1)
    add_table(doc, ["数据类别", "常见物理量", "常见图/表", "数据库用途"], COMMON_DATA)
    doc.add_heading("3. 常见结构与研究内容矩阵", level=1)
    add_table(doc, ["结构类型", "典型结构", "主要研究方面", "常见可入库数据"], STRUCTURES)
    doc.add_heading("4. 入库优先级建议", level=1)
    for item in [
        "优先录入结构参数、材料、光谱坐标、Q、lambda/frequency、linewidth/Fano q。",
        "随后录入 k-space、band structure、polarization vortex/topological charge，用来判断机制和拓扑性质。",
        "near-field、multipole、LDOS/PLDOS、field enhancement 用于解释增强机制和应用性能。",
        "所有数值必须保留 source_quote、page/figure/table、verification_status；图谱数字化必须记录不确定度和拟合公式。",
    ]:
        doc.add_paragraph(item, style="List Bullet")
    doc.add_heading("5. 数据库当前状态", level=1)
    doc.add_paragraph("当前数据库中的 physical_observations 表已经包含本地 PDF 和联网网页的 source-quoted 候选物理观测。多数记录仍需人工审阅后才能升级为最终事实。")
    out.parent.mkdir(parents=True, exist_ok=True)
    doc.save(out)


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--db", type=Path, default=DEFAULT_DB)
    parser.add_argument("--docx", type=Path, default=DEFAULT_DOCX)
    parser.add_argument("--md", type=Path, default=DEFAULT_MD)
    args = parser.parse_args()
    conn = sqlite3.connect(args.db)
    text = build_text(conn)
    conn.close()
    args.md.write_text(text, encoding="utf-8")
    build_docx(text, args.docx)
    print(f"Wrote {args.md}")
    print(f"Wrote {args.docx}")


if __name__ == "__main__":
    main()
