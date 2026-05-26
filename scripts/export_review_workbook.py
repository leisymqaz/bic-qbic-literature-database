#!/usr/bin/env python
"""Export the verified BIC database to an Excel workbook for review."""

from __future__ import annotations

import argparse
import sqlite3
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE


PROJECT_DIR = Path(__file__).resolve().parents[1]
DEFAULT_DB = PROJECT_DIR / "data" / "bic_literature_verified.sqlite"
DEFAULT_OUT = PROJECT_DIR / "outputs" / "BIC_verified_database_review.xlsx"


SHEETS = {
    "README": None,
    "Papers": "SELECT * FROM papers ORDER BY paper_id",
    "Physical_Observations": "SELECT * FROM physical_observations ORDER BY source_kind, paper_id, external_id, observation_type",
    "Core_Verified": """
        SELECT p.paper_id, sf.file_name, p.title, p.year, p.source, p.doi, p.arxiv, p.verification_status
        FROM papers p JOIN source_files sf USING(file_id)
        WHERE p.verification_status='manual_pdf_verified_core'
        ORDER BY p.year, p.title
    """,
    "External_Sources": "SELECT * FROM external_verified_papers ORDER BY year, title",
    "Web_Sources": "SELECT * FROM web_sources ORDER BY year, title",
    "Evidence": "SELECT * FROM evidence ORDER BY evidence_id",
    "Candidates": "SELECT * FROM candidate_classifications ORDER BY paper_id, candidate_field",
}


def rows(conn: sqlite3.Connection, query: str) -> tuple[list[str], list[sqlite3.Row]]:
    cur = conn.execute(query)
    return [d[0] for d in cur.description], cur.fetchall()


def write_table(ws, headers: list[str], data: list[sqlite3.Row]) -> None:
    ws.append(headers)
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1F4E78")
        cell.alignment = Alignment(wrap_text=True, vertical="center")
    for row in data:
        ws.append([safe_cell(row[h]) for h in headers])
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    for col_idx, header in enumerate(headers, start=1):
        values = [str(header)] + [str(row[header] or "") for row in data[:300]]
        width = min(max(max(len(v) for v in values) + 2, 10), 60)
        ws.column_dimensions[get_column_letter(col_idx)].width = width
    for row in ws.iter_rows():
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")


def safe_cell(value):
    if value is None:
        return None
    if isinstance(value, (int, float, bool)):
        return value
    text = str(value)
    text = ILLEGAL_CHARACTERS_RE.sub("", text)
    if len(text) > 32000:
        text = text[:32000] + " ...[truncated]"
    return text


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--db", type=Path, default=DEFAULT_DB)
    parser.add_argument("--out", type=Path, default=DEFAULT_OUT)
    args = parser.parse_args()
    conn = sqlite3.connect(args.db)
    conn.row_factory = sqlite3.Row

    wb = Workbook()
    default = wb.active
    wb.remove(default)

    readme = wb.create_sheet("README")
    readme_rows = [
        ["Workbook purpose", "Reviewable export of the BIC/qBIC verified database."],
        ["Important boundary", "Physical_Observations rows are source-quoted review queues unless verification_status says human/manual verified."],
        ["Do not treat as facts", "Candidate classifications, regex extracted Q/lambda/Fano values, and any row marked candidate_needs_review."],
        ["Primary review workflow", "Filter Physical_Observations by observation_type, inspect quote/source/page, then promote or reject in SQLite."],
        ["SQLite source", str(args.db)],
    ]
    write_table(readme, ["Field", "Value"], [dict(zip(["Field", "Value"], r)) for r in readme_rows])

    for sheet_name, query in SHEETS.items():
        if sheet_name == "README":
            continue
        ws = wb.create_sheet(sheet_name[:31])
        headers, data = rows(conn, query)
        write_table(ws, headers, data)

    args.out.parent.mkdir(parents=True, exist_ok=True)
    wb.save(args.out)
    conn.close()
    print(f"Wrote workbook: {args.out}")


if __name__ == "__main__":
    main()
